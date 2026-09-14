import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_resume():
    doc = docx.Document()

    # Page setup - Margins (0.5 in top/bottom, 0.55 in left/right)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    # Colors
    NAVY = RGBColor(27, 54, 93)      # #1B365D Primary accent
    SLATE = RGBColor(70, 80, 95)     # #46505F Subheadings/dates
    CHARCOAL = RGBColor(35, 35, 35)  # #232323 Body text

    def set_font(run, name="Calibri", size_pt=10, bold=False, italic=False, color=CHARCOAL):
        run.font.name = name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(title.upper())
        set_font(run, name="Calibri", size_pt=11, bold=True, color=NAVY)
        
        # Add bottom border under paragraph using oxml
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="1B365D"/></w:pBdr>')
        pPr.append(pBdr)

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    
    r_name = p_name.add_run("CHASE KINSLOW, J.D., CPA")
    set_font(r_name, name="Calibri", size_pt=18, bold=True, color=NAVY)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(6)
    
    r_contact = p_contact.add_run("Monroe, LA 71201  |  (501) 707-7779  |  chasekn30@yahoo.com  |  linkedin.com/in/chasekinslow")
    set_font(r_contact, name="Calibri", size_pt=9.5, bold=False, color=SLATE)

    # 2. PROFESSIONAL SUMMARY
    add_section_header("Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(2)
    p_sum.paragraph_format.space_after = Pt(4)
    p_sum.paragraph_format.line_spacing = 1.15
    r_sum = p_sum.add_run(
        "Detail-driven Legal Operations and Compliance Professional with a Juris Doctor, active Arkansas Bar license, and CPA credentials. "
        "Proven background in complex regulatory interpretation, statutory compliance, audit-grade recordkeeping, and cross-functional workflow management. "
        "Highly skilled in translating complex legal frameworks (CCPA, GDPR, TCPA, ADA, APA) into operational reality, managing Data Processing Addenda (DPAs) "
        "and privacy notices, and leveraging GenAI prompt engineering and legal tech to streamline multi-jurisdictional compliance workflows."
    )
    set_font(r_sum, name="Calibri", size_pt=9.5, color=CHARCOAL)

    # 3. CORE COMPETENCIES (2-column neat table)
    add_section_header("Core Competencies")
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(3.7), Inches(3.7)]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]

    competencies = [
        ("• Privacy Program Administration (CCPA, GDPR, TCPA)", "• Legal Tech, GenAI Systems & Prompt Engineering"),
        ("• Data Processing Addenda (DPAs) & Privacy Terms", "• Cross-Functional Alignment (Product, Eng, Security)"),
        ("• Product Legal Compliance & Feature Documentation", "• Standard Operating Procedures (SOP) & Audit Readiness"),
        ("• Regulatory Research, Statutory Analysis & Rulemaking", "• Consumer Protection, Marketing Laws & ADA Accessibility")
    ]

    for row_idx, (c1, c2) in enumerate(competencies):
        cell1 = table.rows[row_idx].cells[0]
        p1 = cell1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(1)
        p1.paragraph_format.space_after = Pt(1)
        r1 = p1.add_run(c1)
        set_font(r1, name="Calibri", size_pt=9, bold=False, color=CHARCOAL)

        cell2 = table.rows[row_idx].cells[1]
        p2 = cell2.paragraphs[0]
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run(c2)
        set_font(r2, name="Calibri", size_pt=9, bold=False, color=CHARCOAL)

    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
            tcPr.append(tcBorders)

    # 4. EDUCATION & CREDENTIALS
    add_section_header("Education & Credentials")
    
    edu_items = [
        ("Juris Doctor (J.D.)", "University of Arkansas School of Law", "Arkansas Bar License"),
        ("Bachelor of Accountancy", "University of Mississippi", "Licensed CPA (AR & LA)")
    ]
    for degree, school, license_text in edu_items:
        p_edu = doc.add_paragraph()
        p_edu.paragraph_format.space_before = Pt(1)
        p_edu.paragraph_format.space_after = Pt(2)
        
        r_deg = p_edu.add_run(f"• {degree}")
        set_font(r_deg, name="Calibri", size_pt=9.5, bold=True, color=CHARCOAL)
        
        r_sch = p_edu.add_run(f" | {school} – ")
        set_font(r_sch, name="Calibri", size_pt=9.5, bold=False, color=CHARCOAL)
        
        r_lic = p_edu.add_run(license_text)
        set_font(r_lic, name="Calibri", size_pt=9.5, bold=True, color=NAVY)

    # 5. PROFESSIONAL EXPERIENCE
    add_section_header("Professional Experience")

    def add_job(title, company, dates, location, bullets):
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(4)
        p_head.paragraph_format.space_after = Pt(1)
        p_head.paragraph_format.keep_with_next = True

        r_title = p_head.add_run(title)
        set_font(r_title, name="Calibri", size_pt=10, bold=True, color=CHARCOAL)

        r_sep = p_head.add_run(" | ")
        set_font(r_sep, name="Calibri", size_pt=10, bold=False, color=SLATE)

        r_co = p_head.add_run(company)
        set_font(r_co, name="Calibri", size_pt=10, bold=True, color=NAVY)

        r_loc = p_head.add_run(f" – {location}")
        set_font(r_loc, name="Calibri", size_pt=9.5, italic=True, color=SLATE)

        p_date = doc.add_paragraph()
        p_date.paragraph_format.space_before = Pt(0)
        p_date.paragraph_format.space_after = Pt(2)
        p_date.paragraph_format.keep_with_next = True
        r_d = p_date.add_run(dates)
        set_font(r_d, name="Calibri", size_pt=9, italic=True, color=SLATE)

        for b in bullets:
            p_b = doc.add_paragraph(style='List Bullet')
            p_b.paragraph_format.space_before = Pt(0)
            p_b.paragraph_format.space_after = Pt(2)
            p_b.paragraph_format.line_spacing = 1.12
            p_b.paragraph_format.left_indent = Inches(0.2)
            r_b = p_b.add_run(b)
            set_font(r_b, name="Calibri", size_pt=9.5, color=CHARCOAL)

    # Job 1: Walmart
    add_job(
        title="Operations & Workflow Specialist",
        company="Walmart",
        dates="May 2026 – Present",
        location="Monroe, LA",
        bullets=[
            "Manage daily cross-functional operational workflows, establishing standard operating procedures (SOPs) to ensure strict adherence to internal compliance policies and operational benchmarks.",
            "Triage business requests and troubleshoot process roadblocks across multi-department teams, maintaining continuous recordkeeping integrity and system alignment.",
            "Integrate modern digital tools and AI-driven process tracking to eliminate operational friction, improve turnaround times, and maintain zero-error execution."
        ]
    )

    # Job 2: Little and Associates
    add_job(
        title="Legal Operations & Regulatory Compliance Specialist",
        company="Little and Associates, LLC",
        dates="December 2020 – December 2023",
        location="Monroe, LA",
        bullets=[
            "Co-administered end-to-end compliance for high-stakes federal and state regulatory programs (IRC § 42, HUD, Fair Housing, ADA), ensuring 100% adherence to statutory standards and agency mandates.",
            "Built complex, audit-ready data models in Excel to maintain comprehensive 'shadow' records and reconcile multi-million-dollar datasets, ensuring complete data integrity for state agency reviews.",
            "Drafted, organized, and reviewed technical compliance exhibits, state filings, and supplemental documentation under strict deadlines with zero audit findings.",
            "Synthesized complex regulatory guidelines into actionable internal checklists and reporting frameworks for cross-functional stakeholders."
        ]
    )

    # Job 3: ADFA
    add_job(
        title="Staff Attorney / Legal Operations Specialist",
        company="Arkansas Development Finance Authority",
        dates="December 2017 – July 2018",
        location="Little Rock, AR",
        bullets=[
            "Optimized state agency legal workflows by internalizing the Administrative Procedure Act (APA) rulemaking process, accelerating regulatory filing timelines and reducing external legal expenditures.",
            "Drafted, reviewed, and maintained official agency rules, public notices, and procedural documentation to ensure full compliance with state and federal statutory requirements.",
            "Directed the state Farm Mediation Program, conducting legal intake, triaging multi-party disputes, and negotiating structured resolutions among lenders, borrowers, and state regulators.",
            "Partnered with executive leadership and cross-functional teams to translate complex statutory updates into operational guidelines and public-facing communications."
        ]
    )

    # Job 4: Murphy and Associates
    add_job(
        title="Client Advisory & Compliance Associate",
        company="Murphy and Associates",
        dates="September 2018 – June 2019",
        location="Monroe, LA",
        bullets=[
            "Managed full-cycle compliance, payroll, and advisory workstreams for a diverse corporate client portfolio, maintaining 100% data integrity and strict regulatory adherence.",
            "Leveraged software automation tools to streamline document intake, reconcile discrepancies, and accelerate complex filing turnaround times."
        ]
    )

    # 6. TECHNICAL & LEGAL TECH SKILLS
    add_section_header("Technical & Legal Tech Skills")
    
    skills = [
        ("Legal & Compliance Tools: ", "Contract Management, DPA Tracking, Regulatory Filing Repositories, AdvanceFlow, UltraTax CS."),
        ("AI & Automation: ", "GenAI Prompt Engineering, Workflow Automation, AI Tool Evaluation, Process Optimization."),
        ("Productivity & Collaboration: ", "Slack, Notion, Microsoft Office 365 (Advanced Excel modeling, Word, PowerPoint), Google Workspace.")
    ]

    for category, detail in skills:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(1)
        p_s.paragraph_format.space_after = Pt(2)
        
        r_cat = p_s.add_run(f"• {category}")
        set_font(r_cat, name="Calibri", size_pt=9.5, bold=True, color=CHARCOAL)
        
        r_det = p_s.add_run(detail)
        set_font(r_det, name="Calibri", size_pt=9.5, bold=False, color=CHARCOAL)

    # Output path
    output_path = r"c:\Users\Charwiz43\Documents\Chase - Personal\Documents\GitHub\chaseknJD.github.io\Chase_Kinslow_Product_Privacy_Paralegal_Resume.docx"
    doc.save(output_path)
    print(f"Successfully generated resume at: {output_path}")

if __name__ == "__main__":
    create_resume()
